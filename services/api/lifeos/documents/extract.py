import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".csv"}
MAX_EXTRACTED_CHARS = 5_000_000
MAX_ARCHIVE_MEMBERS = 500
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_SOURCE_SEGMENTS = 1000
MAX_WORKSHEET_ROWS = 100_000
CHUNK_CHARS = 1800


class DocumentExtractionError(ValueError):
    """A safe, user-displayable document validation or extraction failure."""


@dataclass(frozen=True)
class ExtractedChunk:
    locator: str
    content: str


@dataclass(frozen=True)
class ExtractedDocument:
    chunks: list[ExtractedChunk]
    source_count: int


def normalized_extension(filename: str) -> str:
    extension = PurePath(filename.replace("\\", "/")).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError("Supported formats: PDF, DOCX, PPTX, XLSX, TXT, MD, and CSV.")
    return extension


def _verify_container(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise DocumentExtractionError("This Office file contains too many internal files to process safely.")
            if sum(member.file_size for member in members) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise DocumentExtractionError("This Office file expands beyond the supported processing limit.")
            if any(member.flag_bits & 0x1 for member in members):
                raise DocumentExtractionError("Password-protected Office files are not supported.")
    except zipfile.BadZipFile as exc:
        raise DocumentExtractionError("This Office file is damaged or has an invalid format.") from exc


def _split_text(locator: str, text: str) -> list[ExtractedChunk]:
    text = re.sub(r"\r\n?", "\n", text)
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        return []
    chunks: list[ExtractedChunk] = []
    current = ""
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        while len(paragraph) > CHUNK_CHARS:
            if current:
                chunks.append(ExtractedChunk(locator, current))
                current = ""
            chunks.append(ExtractedChunk(locator, paragraph[:CHUNK_CHARS]))
            paragraph = paragraph[CHUNK_CHARS:]
        if current and len(current) + len(paragraph) + 2 > CHUNK_CHARS:
            chunks.append(ExtractedChunk(locator, current))
            current = paragraph
        else:
            current = f"{current}\n\n{paragraph}" if current else paragraph
    if current:
        chunks.append(ExtractedChunk(locator, current))
    return chunks


def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    extension = normalized_extension(filename)
    if not data:
        raise DocumentExtractionError("The uploaded file is empty.")

    segments: list[tuple[str, str]] = []
    if extension == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise DocumentExtractionError("The file content does not match a PDF document.")
        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise DocumentExtractionError("Password-protected PDFs are not supported.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise DocumentExtractionError(f"PDFs are limited to {MAX_PDF_PAGES} pages.")
            segments = [(f"Page {index + 1}", page.extract_text() or "") for index, page in enumerate(reader.pages)]
        except DocumentExtractionError:
            raise
        except Exception as exc:
            raise DocumentExtractionError("This PDF could not be read. Check that it is a valid, unprotected PDF.") from exc
    elif extension in {".docx", ".pptx", ".xlsx"}:
        _verify_container(data)
        try:
            if extension == ".docx":
                doc = WordDocument(io.BytesIO(data))
                segments = [("Document", "\n\n".join(p.text for p in doc.paragraphs))]
                for table_index, table in enumerate(doc.tables, 1):
                    rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
                    segments.append((f"Table {table_index}", "\n".join(rows)))
            elif extension == ".pptx":
                presentation = Presentation(io.BytesIO(data))
                if len(presentation.slides) > MAX_SOURCE_SEGMENTS:
                    raise DocumentExtractionError("This presentation contains too many slides to process safely.")
                segments = [
                    (f"Slide {index}", "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)))
                    for index, slide in enumerate(presentation.slides, 1)
                ]
            else:
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                if len(workbook.worksheets) > MAX_SOURCE_SEGMENTS:
                    raise DocumentExtractionError("This workbook contains too many sheets to process safely.")
                for sheet in workbook.worksheets:
                    lines = []
                    extracted_chars = 0
                    for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if row_index > MAX_WORKSHEET_ROWS:
                            raise DocumentExtractionError(
                                f"Worksheets are limited to {MAX_WORKSHEET_ROWS:,} rows."
                            )
                        values = [str(value) if value is not None else "" for value in row]
                        if any(values):
                            line = f"Row {row_index}: " + " | ".join(values)
                            extracted_chars += len(line)
                            if extracted_chars > MAX_EXTRACTED_CHARS:
                                raise DocumentExtractionError("Extracted text exceeds the supported processing limit.")
                            lines.append(line)
                    segments.append((f"Sheet {sheet.title}", "\n".join(lines)))
                workbook.close()
        except DocumentExtractionError:
            raise
        except Exception as exc:
            raise DocumentExtractionError("This Office file could not be read. Check that it is valid and unprotected.") from exc
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentExtractionError("Text and CSV files must use UTF-8 encoding.") from exc
        if extension == ".csv":
            try:
                rows = list(csv.reader(io.StringIO(text), strict=True))
            except csv.Error as exc:
                raise DocumentExtractionError("This CSV file contains invalid rows.") from exc
            text = "\n".join(f"Row {index}: " + " | ".join(row) for index, row in enumerate(rows, 1))
        segments = [("Document", text)]

    if len(segments) > MAX_SOURCE_SEGMENTS:
        raise DocumentExtractionError("This file contains too many sections to process safely.")
    if sum(len(text) for _, text in segments) > MAX_EXTRACTED_CHARS:
        raise DocumentExtractionError("Extracted text exceeds the supported processing limit.")
    chunks = [chunk for locator, text in segments for chunk in _split_text(locator, text)]
    if not chunks:
        raise DocumentExtractionError(
            "No readable text was found. Scanned PDFs and image-only documents need OCR, which is not enabled."
        )
    return ExtractedDocument(chunks=chunks, source_count=len(segments))
