import uuid
from io import BytesIO

import pytest
from docx import Document as WordDocument
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from lifeos.core.config import settings
from lifeos.core.database import init_db
from lifeos.documents.extract import extract_document
from lifeos.main import app


async def _register(client: AsyncClient) -> dict[str, str]:
    response = await client.post(
        "/api/auth/register",
        json={
            "email": f"documents-{uuid.uuid4().hex}@convia.app",
            "password": "test-pass-123",
            "full_name": "Document Test",
        },
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_document_extraction_retains_source_locations():
    word = WordDocument()
    word.add_paragraph("A useful project note with a distinctive phrase.")
    stream = BytesIO()
    word.save(stream)

    extracted = extract_document("notes.docx", stream.getvalue())

    assert extracted.source_count == 1
    assert extracted.chunks[0].locator == "Document"
    assert "distinctive phrase" in extracted.chunks[0].content


def test_document_extraction_rejects_extension_content_mismatch():
    with pytest.raises(ValueError, match="does not match a PDF"):
        extract_document("notes.pdf", b"not really a pdf")


def test_document_extraction_supports_csv_xlsx_and_pptx():
    csv_doc = extract_document("budget.csv", b"category,amount\ntravel,4200\n")
    assert csv_doc.chunks[0].locator == "Document"
    assert "Row 2: travel | 4200" in csv_doc.chunks[0].content

    workbook = Workbook()
    workbook.active["A1"] = "budget"
    workbook.active["B1"] = 4200
    xlsx_stream = BytesIO()
    workbook.save(xlsx_stream)
    xlsx_doc = extract_document("budget.xlsx", xlsx_stream.getvalue())
    assert xlsx_doc.chunks[0].locator == "Sheet Sheet"
    assert "budget | 4200" in xlsx_doc.chunks[0].content

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Quarterly review"
    pptx_stream = BytesIO()
    presentation.save(pptx_stream)
    pptx_doc = extract_document("review.pptx", pptx_stream.getvalue())
    assert pptx_doc.chunks[0].locator == "Slide 1"
    assert "Quarterly review" in pptx_doc.chunks[0].content


@pytest.mark.asyncio
async def test_document_upload_is_private_searchable_and_deletable(tmp_path, monkeypatch):
    await init_db()
    monkeypatch.setattr(settings, "DOCUMENT_STORAGE_DIR", str(tmp_path / "private-documents"))
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        owner_headers = await _register(client)
        other_headers = await _register(client)
        upload = await client.post(
            "/api/documents",
            headers=owner_headers,
            files={"file": ("project-notes.txt", b"Budget is 4200 and the delivery deadline is Friday.", "text/plain")},
        )
        assert upload.status_code == 201, upload.text
        document = upload.json()
        assert document["chunk_count"] == 1
        assert document["source_count"] == 1
        stored_files = list((tmp_path / "private-documents").rglob(f"{document['id']}.txt"))
        assert len(stored_files) == 1

        owner_list = await client.get("/api/documents", headers=owner_headers)
        other_list = await client.get("/api/documents", headers=other_headers)
        assert [row["id"] for row in owner_list.json()] == [document["id"]]
        assert other_list.json() == []
        hidden_detail = await client.get(f"/api/documents/{document['id']}", headers=other_headers)
        assert hidden_detail.status_code == 404

        answer = await client.post(
            f"/api/documents/{document['id']}/ask",
            headers=owner_headers,
            json={"question": "What is the delivery deadline?"},
        )
        assert answer.status_code == 200
        assert answer.json()["status"] == "AI_NOT_CONFIGURED"
        assert answer.json()["answer"] is None
        assert answer.json()["sources"][0]["locator"] == "Document"
        assert "Friday" in answer.json()["sources"][0]["excerpt"]

        preview = await client.get(f"/api/documents/{document['id']}/preview", headers=owner_headers)
        assert preview.status_code == 200
        assert preview.json()["passages"][0]["locator"] == "Document"
        assert (await client.get(f"/api/documents/{document['id']}/preview", headers=other_headers)).status_code == 404

        summary = await client.post(f"/api/documents/{document['id']}/summary", headers=owner_headers, json={"mode": "action"})
        assert summary.status_code == 200
        assert summary.json()["status"] == "AI_NOT_CONFIGURED"
        assert summary.json()["answer"] is None

        deleted = await client.delete(f"/api/documents/{document['id']}", headers=owner_headers)
        assert deleted.status_code == 204
        assert not stored_files[0].exists()
        assert (await client.get("/api/documents", headers=owner_headers)).json() == []


@pytest.mark.asyncio
async def test_document_answer_rejects_unmapped_citation(tmp_path, monkeypatch):
    await init_db()
    monkeypatch.setattr(settings, "DOCUMENT_STORAGE_DIR", str(tmp_path / "citation-documents"))
    from lifeos.api.documents import nemotron_provider

    async def configured():
        return {"configured": True, "status": "CONFIGURED"}

    async def fabricated_citation(**kwargs):
        return {"content": "The deadline is Friday. [S99]", "provider": "mock", "model_used": "mock"}

    monkeypatch.setattr(nemotron_provider, "health_check", configured)
    monkeypatch.setattr(nemotron_provider, "generate", fabricated_citation)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        headers = await _register(client)
        uploaded = await client.post("/api/documents", headers=headers, files={"file": ("note.txt", b"Deadline is Friday.", "text/plain")})
        document_id = uploaded.json()["id"]
        answer = await client.post(f"/api/documents/{document_id}/ask", headers=headers, json={"question": "When is the deadline?"})
    assert answer.status_code == 200
    assert answer.json()["status"] == "CITATION_VALIDATION_FAILED"
    assert answer.json()["answer"] is None


@pytest.mark.asyncio
async def test_document_upload_requires_authentication():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post(
            "/api/documents",
            files={"file": ("note.txt", b"private text", "text/plain")},
        )
    assert response.status_code == 401


@pytest.mark.asyncio
async def test_document_upload_enforces_account_storage_quota(tmp_path, monkeypatch):
    await init_db()
    monkeypatch.setattr(settings, "DOCUMENT_STORAGE_DIR", str(tmp_path / "quota-documents"))
    monkeypatch.setattr(settings, "DOCUMENT_MAX_USER_BYTES", 20)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers = await _register(client)
        response = await client.post(
            "/api/documents",
            headers=headers,
            files={"file": ("large.txt", b"This document is over the tiny test quota.", "text/plain")},
        )
    assert response.status_code == 413
    assert response.json()["detail"] == "Your document vault has reached its storage limit."
